"""Read-only conversion of the supplied ALas Excel files. Requires openpyxl."""
import json, sys, hashlib
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
source = Path(sys.argv[1]) if len(sys.argv) > 1 else root.parent / 'upload'
municipalities = []
audit = []
for path in sorted(source.glob('ALas_*.xlsx')):
    book = openpyxl.load_workbook(path, data_only=False)
    sheet = book['KHK-päästöt']
    rows = list(sheet.values)
    name = rows[0][0].capitalize()
    years = [int(v) for v in rows[0][1:]]
    sectors = [r[0] for r in rows[1:14]]
    assert len(sectors) == 13 and len(set(years)) == len(years)
    assert rows[14][0] == 'Päästöhyvitykset'
    assert rows[19][0] == 'Hinku-laskenta ilman päästöhyvityksiä'
    assert all(isinstance(v, (int, float)) for r in rows[1:18] for v in r[1:])
    observations = {}
    for j, year in enumerate(years, 1):
        values = {r[0]: r[j] for r in rows[1:14]}
        observations[str(year)] = {'sectors': values, 'total': rows[15][j], 'credits': rows[14][j], 'perCapita': rows[16][j], 'population': rows[17][j]}
    municipalities.append({'id': name.lower(), 'name': name, 'sectors': sectors, 'years': years, 'observations': observations, 'metadata': {'source': 'Suomen ympäristökeskus / Hiilineutraalisuomi', 'url': 'https://paastot.hiilineutraalisuomi.fi/', 'method': rows[19][0], 'unit': 'kt CO₂e', 'fileName': path.name, 'downloadDate': '2026-09-15', 'publicationDate': None, 'importedAt': '2026-09-15'}})
    audit.append({'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'worksheets': [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column, 'mergedCells': len(s.merged_cells.ranges), 'formulas': sum(c.data_type == 'f' for row in s for c in row)} for s in book], 'years': years, 'sectors': sectors, 'roundingDifferenceMaxKt': round(max(abs(o['total'] - sum(o['sectors'].values())) for o in observations.values()), 6)})
assert len(municipalities) == 13
data = {'schemaVersion': 1, 'publishedRevision': 'initial-2026-09-15', 'municipalities': municipalities, 'configs': {}, 'actions': []}
(root / 'data' / 'dataset.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
(root / 'data' / 'source-audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
print('Converted', len(municipalities), 'municipalities; observations', sum(len(m['years']) for m in municipalities))
